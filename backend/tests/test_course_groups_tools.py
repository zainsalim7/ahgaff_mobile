"""Backend tests for 👥 Course Groups tools (iteration 106)."""
import io
import os
import openpyxl
import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
COURSE_ID = "698f0000d803b27aab0120af"
DEPT_ID = "698e501d97fef774e66e93a9"
LEVEL = 3


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": "admin", "password": "admin123"}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def h(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


def test_group_detail(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/1/detail", headers=h, timeout=30)
    assert r.status_code == 200, r.text
    j = r.json()
    for k in ("course", "group", "all_groups", "students", "limits", "schedule", "lectures"):
        assert k in j, f"missing {k}"
    assert j["group"]["key"] == "1"
    assert isinstance(j["students"], list) and len(j["students"]) > 0
    s = j["students"][0]
    assert "pct_unexcused" in s and "pct_total" in s and "status_ar" in s
    assert any(g["key"] == "2" for g in j["all_groups"])


def test_group_detail_not_found(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/99/detail", headers=h, timeout=30)
    assert r.status_code == 404


def test_export_all_groups(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/export", headers=h, timeout=30)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    assert r.headers.get("X-Filename")
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    names = wb.sheetnames
    assert "بلا مجموعة" in names
    assert "الكل" in names
    # at least one group sheet (مجموعة ...)
    assert len(names) >= 3


def test_export_single_group(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/export?group=2", headers=h, timeout=30)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert len(wb.sheetnames) == 1


def test_export_unknown_group(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/export?group=ZZ", headers=h, timeout=30)
    assert r.status_code == 404


def test_export_level(h):
    r = requests.get(f"{BASE_URL}/api/groups/export-level?department_id={DEPT_ID}&level={LEVEL}", headers=h, timeout=60)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames[0] == "ملخص"
    assert len(wb.sheetnames) >= 2


def test_export_level_empty(h):
    r = requests.get(f"{BASE_URL}/api/groups/export-level?department_id={DEPT_ID}&level=99", headers=h, timeout=30)
    assert r.status_code == 404


def test_import_template(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/import-template", headers=h, timeout=30)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "رقم القيد" in header
    assert "اسم الطالب" in header
    assert "المجموعة" in header


def _build_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["رقم القيد", "اسم الطالب", "المجموعة"])
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


def test_import_preview_mixed(h):
    # First get detail to know real students
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/1/detail", headers=h, timeout=30)
    g1 = r.json()
    student_g1 = g1["students"][0]  # currently in group 1
    r2 = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/2/detail", headers=h, timeout=30)
    student_g2 = r2.json()["students"][0]
    # row1: move student from 1 -> 2
    # row2: unchanged (group2 student stays in 2)
    # row3: unknown student 999999
    # row4: known student with undefined group 'ج'
    student_g1b = g1["students"][1]
    rows = [
        [student_g1["student_number"], student_g1["full_name"], "2"],
        [student_g2["student_number"], student_g2["full_name"], "2"],
        ["999999", "طالب غير مسجل", "1"],
        [student_g1b["student_number"], student_g1b["full_name"], "ج"],
    ]
    data = _build_xlsx(rows)
    files = {"file": ("test.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/import-preview", headers=h, files=files, timeout=30)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["unchanged"] >= 1
    assert any(u["student_number"] == "999999" for u in j["unknown_students"])
    assert any(u["label"] == "ج" for u in j["undefined_groups"])
    assert j["changes"] >= 1
    # store for apply test
    pytest.student_to_move = student_g1
    pytest.student_g2_unchanged = student_g2


def test_import_preview_bad_header(h):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["foo", "bar"])
    ws.append(["x", "y"])
    bio = io.BytesIO()
    wb.save(bio)
    files = {"file": ("bad.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/import-preview", headers=h, files=files, timeout=30)
    assert r.status_code == 400


def test_import_apply_rejects_undefined(h):
    student = pytest.student_to_move
    rows = [[student["student_number"], student["full_name"], "ج"]]
    data = _build_xlsx(rows)
    files = {"file": ("x.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/import-apply", headers=h, files=files, timeout=30)
    assert r.status_code == 400


def test_import_apply_valid_and_revert(h):
    student = pytest.student_to_move  # was in group 1
    rows = [[student["student_number"], student["full_name"], "2"]]
    data = _build_xlsx(rows)
    files = {"file": ("x.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/import-apply", headers=h, files=files, timeout=30)
    assert r.status_code == 200, r.text
    assert r.json()["applied"] >= 1
    # Verify via detail: now in group 2
    g2 = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/2/detail", headers=h, timeout=30).json()
    assert any(s["student_number"] == student["student_number"] for s in g2["students"]), "student should be in group 2 now"
    # revert via assign
    sid = next(s["student_id"] for s in g2["students"] if s["student_number"] == student["student_number"])
    r2 = requests.post(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/assign", headers=h,
                       json={"student_ids": [sid], "group": "1"}, timeout=30)
    assert r2.status_code == 200, r2.text
    g1 = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups/1/detail", headers=h, timeout=30).json()
    assert any(s["student_number"] == student["student_number"] for s in g1["students"]), "student reverted"


def test_students_groups_map(h):
    # Use the student 698ef706d803b27aab012066 from req
    sid = "698ef706d803b27aab012066"
    r = requests.post(f"{BASE_URL}/api/students/groups-map", headers=h, json={"student_ids": [sid]}, timeout=30)
    assert r.status_code == 200, r.text
    j = r.json()
    assert isinstance(j, dict)
    if sid in j:
        for row in j[sid]:
            for k in ("course_id", "course_name", "group", "group_name"):
                assert k in row


def test_put_groups_returns_warnings(h):
    r = requests.get(f"{BASE_URL}/api/courses/{COURSE_ID}/groups", headers=h, timeout=30)
    assert r.status_code == 200
    data = r.json()
    # Build identical PUT payload
    groups_payload = [{"key": g["key"], "name": g["name"], "teacher_id": g.get("teacher_id")} for g in data.get("groups", [])]
    r2 = requests.put(f"{BASE_URL}/api/courses/{COURSE_ID}/groups", headers=h, json={"groups": groups_payload}, timeout=30)
    assert r2.status_code == 200, r2.text
    j = r2.json()
    assert "schedule_sync_warnings" in j
    assert isinstance(j["schedule_sync_warnings"], list)
