"""اختبار استيراد خلية متعددة المجموعات (dry-run) على بيئة المعاينة"""
import io
import os
import sys

import openpyxl
import requests

API = [l.split("=", 1)[1].strip() for l in open("/app/frontend/.env") if l.startswith("REACT_APP_BACKEND_URL")][0]
FAC, DEPT = "698e4fb497fef774e66e93a6", "698e501d97fef774e66e93a9"
tok = requests.post(f"{API}/api/auth/login", json={"username": "admin", "password": "admin123"}).json()
TOKEN = tok.get("access_token") or tok.get("token")
H = {"Authorization": f"Bearer {TOKEN}"}

r = requests.get(f"{API}/api/weekly-schedule/import-template", params={"faculty_id": FAC, "department_id": DEPT}, headers=H)
assert r.status_code == 200, r.text[:200]
wb = openpyxl.load_workbook(io.BytesIO(r.content))
ws = wb.worksheets[0]

# إيجاد صف المستوى 3 وعمود أول يوم/فترة
level_row = None
for row in range(1, ws.max_row + 1):
    v = str(ws.cell(row=row, column=1).value or "")
    if "3" in v or "الثالث" in v:
        level_row = row
        break
assert level_row, "لم يُعثر على صف المستوى 3"
col = 3
mode = sys.argv[1] if len(sys.argv) > 1 else "ok"
if mode == "ok":
    ws.cell(row=level_row, column=col).value = "مكتبة - اختبار (A)\nلغة إنجليزية - اختبار (B)"
    ws.cell(row=level_row + 1, column=col).value = "مختبر اختبار 1\nمختبر اختبار 2"
    ws.cell(row=level_row + 2, column=col).value = "Saeed"
elif mode == "arabic":
    ws.cell(row=level_row, column=col).value = "مكتبة - اختبار (أ)؛ لغة إنجليزية - اختبار (ب)"
    ws.cell(row=level_row + 1, column=col).value = "مختبر اختبار 1\nمختبر اختبار 2"
    ws.cell(row=level_row + 2, column=col).value = "Saeed"
elif mode == "badgroup":
    ws.cell(row=level_row, column=col).value = "مكتبة - اختبار (C)\nلغة إنجليزية - اختبار (B)"
    ws.cell(row=level_row + 1, column=col).value = "مختبر اختبار 1\nمختبر اختبار 2"
    ws.cell(row=level_row + 2, column=col).value = "Saeed"
elif mode == "sameroom":
    ws.cell(row=level_row, column=col).value = "مكتبة - اختبار (A)\nلغة إنجليزية - اختبار (B)"
    ws.cell(row=level_row + 1, column=col).value = "مختبر اختبار 1"
    ws.cell(row=level_row + 2, column=col).value = "Saeed"
elif mode == "overlap":  # نفس المجموعة A للمقررين → طلاب مشتركون
    ws.cell(row=level_row, column=col).value = "مكتبة - اختبار (A)\nلغة إنجليزية - اختبار (A)"
    ws.cell(row=level_row + 1, column=col).value = "مختبر اختبار 1\nمختبر اختبار 2"
    ws.cell(row=level_row + 2, column=col).value = "Saeed"
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
r = requests.post(f"{API}/api/weekly-schedule/import-master", data={"faculty_id": FAC, "department_id": DEPT, "dry_run": os.environ.get("DRY", "1")},
                  files={"file": ("t.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=H, timeout=120)
print("status", r.status_code)
rep = r.json()
for k in ("to_create", "to_replace", "to_group", "grouped", "errors", "conflicts", "can_commit"):
    v = rep.get(k)
    if isinstance(v, list):
        v = [x for x in v if "اختبار" in x or "مجموعة" in x or "مجموع" in x][:6]
    print(k, ":", v)
