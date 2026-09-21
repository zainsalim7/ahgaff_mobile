"""Tests for management dashboard endpoints (لوحة القيادة)."""
import os
import pytest
import requests
from urllib.parse import unquote

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://schedule-hub-272.preview.emergentagent.com').rstrip('/')

FAC_SHARIA = "698e4f9297fef774e66e93a4"
FAC_BANAT = "698e4fb497fef774e66e93a6"
DEPT_SHARIA_QANUN = "698e500997fef774e66e93a8"
DEPT_ISLAMIC_STUDIES = "698e501d97fef774e66e93a9"


def _login(username, password):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": username, "password": password}, timeout=30)
    assert r.status_code == 200, f"Login failed for {username}: {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_token():
    return _login("admin", "admin123")


@pytest.fixture(scope="module")
def salim_token():
    return _login("Salim", "test1234")


@pytest.fixture(scope="module")
def saeed_token():
    return _login("Saeed", "test1234")


@pytest.fixture(scope="module")
def teacher_token():
    return _login("9999", "teacher123")


def _get(token, path, params=None):
    return requests.get(f"{BASE_URL}{path}", headers={"Authorization": f"Bearer {token}"}, params=params or {}, timeout=60)


# ── Admin dashboard basics
class TestAdminDashboard:
    def test_admin_week(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management", {"period": "week"})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["scope"]["label"] == "كل الجامعة"
        assert j["scope"]["is_admin"] is True
        assert j["scope"]["can_filter"] is True
        assert isinstance(j["scope"]["faculties"], list) and len(j["scope"]["faculties"]) > 0
        assert isinstance(j["scope"]["departments"], list) and len(j["scope"]["departments"]) > 0
        n = j["numbers"]
        for k in ("students", "teachers", "courses", "lectures_today", "lectures_period", "attendance_rate"):
            assert k in n
        # Expected fixture numbers
        assert n["lectures_period"] == 31, f"lectures_period={n['lectures_period']}"
        assert abs(n["attendance_rate"] - 87.7) < 0.5, f"attendance_rate={n['attendance_rate']}"
        assert j["chart"]["group_by"] == "date"
        assert len(j["chart"]["points"]) == 7
        assert len(j["alerts"]) == 5
        alerts = {a["key"]: a for a in j["alerts"]}
        assert alerts["low_attendance"]["count"] == 2
        assert alerts["missed_today"]["count"] == 1
        assert alerts["late_teachers"]["count"] == 1
        assert alerts["cancelled"]["count"] == 5
        assert alerts["pending_fees"]["count"] == 1
        assert "types" in j["finance"]
        assert isinstance(j["activity"], list) and len(j["activity"]) <= 15

    def test_admin_day(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management", {"period": "day"})
        assert r.status_code == 200
        j = r.json()
        assert j["chart"]["group_by"] == "department"
        assert isinstance(j["chart"]["points"], list)

    def test_admin_month(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management", {"period": "month"})
        assert r.status_code == 200
        j = r.json()
        assert j["chart"]["group_by"] == "date"
        assert len(j["chart"]["points"]) == 30

    def test_invalid_period_falls_back_to_week(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management", {"period": "yearly"})
        assert r.status_code == 200
        j = r.json()
        assert j["period"] == "week"
        assert len(j["chart"]["points"]) == 7


# ── RBAC
class TestRbac:
    def test_dean_salim(self, salim_token):
        r = _get(salim_token, "/api/dashboard/management", {"period": "week"})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["scope"]["label"] == "كلية الشريعة والقانون", j["scope"]["label"]
        assert j["scope"]["can_filter"] is True
        assert len(j["scope"]["departments"]) == 3
        assert j["numbers"]["students"] == 19, j["numbers"]["students"]

    def test_dept_head_saeed(self, saeed_token):
        r = _get(saeed_token, "/api/dashboard/management", {"period": "week"})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["scope"]["label"] == "الدراسات الإسلامية", j["scope"]["label"]
        assert j["scope"]["can_filter"] is False
        assert j["numbers"]["students"] == 72, j["numbers"]["students"]
        assert j["numbers"]["lectures_period"] == 0

    def test_teacher_forbidden(self, teacher_token):
        r = _get(teacher_token, "/api/dashboard/management", {"period": "week"})
        assert r.status_code == 403
        assert "لوحة القيادة" in r.json().get("detail", "")


# ── Manual filter
class TestManualFilter:
    def test_salim_own_dept(self, salim_token):
        r = _get(salim_token, "/api/dashboard/management", {"period": "week", "department_id": DEPT_SHARIA_QANUN})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["scope"]["label"].strip() == "الشريعة والقانون", j["scope"]["label"]

    def test_salim_other_faculty_dept_forbidden(self, salim_token):
        r = _get(salim_token, "/api/dashboard/management", {"period": "week", "department_id": DEPT_ISLAMIC_STUDIES})
        assert r.status_code == 403
        assert "خارج نطاق" in r.json().get("detail", "")

    def test_salim_other_faculty_forbidden(self, salim_token):
        r = _get(salim_token, "/api/dashboard/management", {"period": "week", "faculty_id": FAC_BANAT})
        assert r.status_code == 403
        assert "خارج نطاق" in r.json().get("detail", "")

    def test_admin_filter_faculty(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management", {"period": "week", "faculty_id": FAC_SHARIA})
        assert r.status_code == 200
        j = r.json()
        assert j["scope"]["label"].strip() == "كلية الشريعة والقانون", j["scope"]["label"]
        assert j["numbers"]["students"] == 19


# ── Export
class TestExport:
    def test_pdf_admin(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management/export", {"fmt": "pdf", "period": "week"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/pdf")
        raw = r.headers.get("X-Filename", "")
        decoded = unquote(raw)
        assert decoded.startswith("لوحة القيادة - "), decoded

    def test_excel_admin(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management/export", {"fmt": "excel", "period": "week"})
        assert r.status_code == 200
        assert "spreadsheet" in r.headers["content-type"]
        # verify sheets
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        required = {"الأرقام", "التنبيهات", "الحضور", "المالية", "سجل النشاط"}
        assert required.issubset(set(wb.sheetnames)), f"sheets={wb.sheetnames}"

    def test_export_forbidden_teacher(self, teacher_token):
        r = _get(teacher_token, "/api/dashboard/management/export", {"fmt": "pdf"})
        assert r.status_code == 403
