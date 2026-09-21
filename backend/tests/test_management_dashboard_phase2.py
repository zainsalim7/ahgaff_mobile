"""Tests for Management Dashboard Phase 2 sections: teachers, students, rooms."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://schedule-hub-272.preview.emergentagent.com').rstrip('/')

DEPT_ISLAMIC_STUDIES = "698e501d97fef774e66e93a9"


def _login(u, p):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"username": u, "password": p}, timeout=30)
    assert r.status_code == 200, f"login failed for {u}: {r.status_code} {r.text[:200]}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_token():
    return _login("admin", "admin123")


@pytest.fixture(scope="module")
def salim_token():
    return _login("Salim", "test1234")


@pytest.fixture(scope="module")
def saeed_token():
    return _login("Saeed", "test1234")


@pytest.fixture(scope="module")
def president_token():
    return _login("president", "test1234")


@pytest.fixture(scope="module")
def teacher_token():
    return _login("9999", "teacher123")


def _get(token, path, params=None):
    return requests.get(f"{BASE_URL}{path}", headers={"Authorization": f"Bearer {token}"}, params=params or {}, timeout=60)


# ── Admin: verify phase 2 structure and expected fixture values
class TestAdminPhase2:
    def test_admin_week_has_phase2_keys(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management", {"period": "week"})
        assert r.status_code == 200, r.text
        j = r.json()
        # keys
        assert "teachers" in j and "students" in j and "rooms" in j
        t = j["teachers"]; s = j["students"]; rm = j["rooms"]
        for k in ("count", "total_hours", "avg_commitment", "avg_delay", "top", "bottom", "rows"):
            assert k in t, f"teachers missing {k}"
        for k in ("by_department", "by_level", "by_status", "warned", "deprived", "evaluated"):
            assert k in s, f"students missing {k}"
        for k in ("capacity_per_room", "rooms_used", "total_slots", "avg_occupancy",
                  "most_used", "least_used", "rows", "unscheduled_courses",
                  "unscheduled_count", "today_heatmap"):
            assert k in rm, f"rooms missing {k}"

    def test_admin_teachers_fixture(self, admin_token):
        j = _get(admin_token, "/api/dashboard/management", {"period": "week"}).json()
        t = j["teachers"]
        assert t["count"] == 2, f"teachers.count={t['count']}"
        assert isinstance(t["rows"], list) and len(t["rows"]) == 2
        # commitment 83.3 expected in at least one row
        commitments = [row.get("commitment") for row in t["rows"]]
        assert any(abs((c or 0) - 83.3) < 0.5 for c in commitments), f"commitments={commitments}"

    def test_admin_students_fixture(self, admin_token):
        j = _get(admin_token, "/api/dashboard/management", {"period": "week"}).json()
        s = j["students"]
        depts = {d["name"].strip(): d for d in s["by_department"]}
        assert len(s["by_department"]) == 2, s["by_department"]
        assert "الدراسات الإسلامية" in depts
        assert "الشريعة والقانون" in depts
        assert depts["الدراسات الإسلامية"].get("students") == 72
        assert depts["الشريعة والقانون"].get("students") == 19
        rate = depts["الشريعة والقانون"].get("rate")
        assert rate is not None and abs(rate - 87.7) < 0.5, f"rate={rate}"
        # by_status includes active 91
        status = {x["status"]: x for x in s["by_status"]}
        assert "active" in status and status["active"].get("count") == 91, s["by_status"]
        # warned/deprived counts (test data)
        assert isinstance(s["warned"], int)
        assert isinstance(s["deprived"], int)

    def test_admin_rooms_fixture(self, admin_token):
        j = _get(admin_token, "/api/dashboard/management", {"period": "week"}).json()
        rm = j["rooms"]
        assert rm["rooms_used"] == 3, rm["rooms_used"]
        assert rm["capacity_per_room"] == 30, rm["capacity_per_room"]
        assert rm["unscheduled_count"] == 4, rm["unscheduled_count"]
        assert isinstance(rm["unscheduled_courses"], list) and len(rm["unscheduled_courses"]) == 4
        assert isinstance(rm["today_heatmap"], list)


# ── Period variants
class TestPeriods:
    @pytest.mark.parametrize("period", ["day", "month"])
    def test_period_has_phase2(self, admin_token, period):
        r = _get(admin_token, "/api/dashboard/management", {"period": period})
        assert r.status_code == 200, r.text
        j = r.json()
        assert "teachers" in j and "students" in j and "rooms" in j


# ── RBAC / scope for phase 2
class TestScopePhase2:
    def test_salim_scope(self, salim_token):
        j = _get(salim_token, "/api/dashboard/management", {"period": "week"}).json()
        dept_names = {d["name"].strip() for d in j["students"]["by_department"]}
        assert "الدراسات الإسلامية" not in dept_names, dept_names
        # rooms slots scoped => total_slots is int
        assert isinstance(j["rooms"]["total_slots"], int)

    def test_salim_invalid_dept_forbidden(self, salim_token):
        r = _get(salim_token, "/api/dashboard/management",
                 {"period": "week", "department_id": DEPT_ISLAMIC_STUDIES})
        assert r.status_code == 403

    def test_saeed_scope(self, saeed_token):
        j = _get(saeed_token, "/api/dashboard/management", {"period": "week"}).json()
        depts = j["students"]["by_department"]
        assert len(depts) == 1, depts
        assert depts[0]["name"].strip() == "الدراسات الإسلامية"
        assert depts[0].get("students") == 72
        assert j["teachers"]["count"] == 0

    def test_president_same_as_admin(self, president_token, admin_token):
        pa = _get(president_token, "/api/dashboard/management", {"period": "week"}).json()
        ad = _get(admin_token, "/api/dashboard/management", {"period": "week"}).json()
        assert pa["teachers"]["count"] == ad["teachers"]["count"]
        assert pa["rooms"]["rooms_used"] == ad["rooms"]["rooms_used"]
        assert sorted([(x["status"], x["count"]) for x in pa["students"]["by_status"]]) == \
               sorted([(x["status"], x["count"]) for x in ad["students"]["by_status"]])

    def test_teacher_forbidden(self, teacher_token):
        r = _get(teacher_token, "/api/dashboard/management", {"period": "week"})
        assert r.status_code == 403


# ── Export phase 2
class TestExportPhase2:
    def test_excel_has_phase2_sheets(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management/export", {"fmt": "excel", "period": "week"})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        required = {"الأساتذة", "الطلاب حسب القسم", "حالات الطلاب", "إشغال القاعات", "مقررات غير مدرجة"}
        missing = required - set(wb.sheetnames)
        assert not missing, f"missing sheets: {missing}; got: {wb.sheetnames}"

    def test_pdf_admin_multipage(self, admin_token):
        r = _get(admin_token, "/api/dashboard/management/export", {"fmt": "pdf", "period": "week"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/pdf")
        # valid PDF starts with %PDF
        assert r.content[:4] == b"%PDF"
        # count pages roughly
        pages = r.content.count(b"/Type /Page") + r.content.count(b"/Type/Page")
        assert pages >= 3, f"pages ~= {pages}"
