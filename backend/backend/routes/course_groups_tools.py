"""👥 المجموعات — صفحة المجموعة، تصدير/استيراد Excel، تصدير المستوى، خريطة مجموعات الطلاب"""
import io
from typing import List, Optional
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from bson import ObjectId
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from .deps import get_db, get_current_user, log_activity
from .course_groups import _can, _course, _view, _oid
from services.absence_engine import course_profiles, course_limits

router = APIRouter()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD_FILL = PatternFill("solid", fgColor="1A2540")
HEAD_FONT = Font(bold=True, color="FFFFFF")


async def teacher_group_keys(db, course: dict, user: dict) -> Optional[List[str]]:
    """👥 None = يرى المقرر كاملاً (إدارة أو منسّق لا يقود مجموعة)؛ قائمة = مجموعاته فقط؛ 403 إن لم يكن له صلة بالمقرر"""
    if user.get("role") != "teacher":
        return None
    u = await db.users.find_one({"_id": ObjectId(user["id"])}, {"teacher_record_id": 1})
    ids = {user["id"], str((u or {}).get("teacher_record_id") or "")} - {""}
    keys = [g.get("key") for g in (course.get("groups") or []) if g.get("teacher_id") in ids and g.get("key")]
    if keys:
        return keys
    if str(course.get("teacher_id")) in ids:
        return None
    raise HTTPException(status_code=403, detail="هذا المقرر ليس ضمن مقرراتك")


def restrict_view(view: dict, keys: Optional[List[str]]) -> dict:
    if keys is None:
        return view
    view["groups"] = [g for g in view["groups"] if g["key"] in keys]
    view["unassigned"] = []
    view["total"] = sum(g["count"] for g in view["groups"])
    view["restricted_to"] = keys
    return view


def _xlsx_response(wb, name: str) -> StreamingResponse:
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fn = quote(f"{name}.xlsx")
    return StreamingResponse(out, media_type=XLSX, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fn}", "X-Filename": fn})


def _sheet(wb, title: str, headers: List[str], rows: List[list], first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31].replace("/", "-").replace("\\", "-").replace("?", "").replace("*", "").replace("[", "(").replace("]", ")").replace(":", "-") or "ورقة"
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, max([len(str(h))] + [len(str(r[i - 1])) for r in rows if len(r) >= i]) + 4))
    return ws


def _group_rows(view: dict) -> List[list]:
    rows = []
    for g in view["groups"]:
        for s in g["students"]:
            rows.append([s["student_number"], s["full_name"], g["name"], g["key"], g.get("teacher_name") or ""])
    for s in view["unassigned"]:
        rows.append([s["student_number"], s["full_name"], "", "", ""])
    return rows


GROUP_HEADERS = ["رقم القيد", "اسم الطالب", "المجموعة", "رمز المجموعة", "قائد المجموعة"]


# ==================== صفحة المجموعة ====================

@router.get("/courses/{course_id}/groups/{key}/detail")
async def group_detail(course_id: str, key: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    course = await _course(db, course_id)
    keys = await teacher_group_keys(db, course, current_user)
    if keys is not None and key not in keys:
        raise HTTPException(status_code=403, detail="هذه المجموعة ليست ضمن مجموعاتك")
    view = restrict_view(await _view(db, course), keys)
    g = next((x for x in view["groups"] if x["key"] == key), None)
    if not g:
        raise HTTPException(status_code=404, detail="المجموعة غير معرّفة في هذا المقرر")
    sids = [s["student_id"] for s in g["students"]]
    profiles = await course_profiles(db, course, sids) if sids else {}
    st_docs = {str(s["_id"]): s for s in await db.students.find({"_id": {"$in": [ObjectId(i) for i in sids if ObjectId.is_valid(i)]}}, {"level": 1, "section": 1, "phone": 1, "is_active": 1}).to_list(2000)} if sids else {}
    students = []
    for s in g["students"]:
        p = profiles.get(s["student_id"], {})
        d = st_docs.get(s["student_id"], {})
        students.append({**s, "level": d.get("level"), "section": d.get("section"), "is_active": d.get("is_active", True),
                         "pct_unexcused": p.get("pct_unexcused", 0), "pct_total": p.get("pct_total", 0), "status": p.get("status", "ok"), "status_ar": p.get("status_ar", "")})
    cells = await db.weekly_schedule.find({"course_id": course_id, "group": key}, {"day": 1, "slot_number": 1, "start_time": 1, "end_time": 1, "room_id": 1, "slot_type": 1}).to_list(100)
    rooms = {str(r["_id"]): r.get("name", "") for r in await db.rooms.find({"_id": {"$in": [ObjectId(c["room_id"]) for c in cells if c.get("room_id") and ObjectId.is_valid(c["room_id"])]}}, {"name": 1}).to_list(100)} if cells else {}
    lec_stats = {}
    async for r in db.lectures.aggregate([{"$match": {"course_id": course_id, "group": key}}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}]):
        lec_stats[r["_id"]] = r["n"]
    dept = await db.departments.find_one({"_id": ObjectId(course["department_id"])}, {"name": 1}) if course.get("department_id") and ObjectId.is_valid(course["department_id"]) else None
    return {
        "course": {"id": course_id, "name": course.get("name", ""), "code": course.get("code", ""), "level": course.get("level"), "section": course.get("section") or "",
                   "department_id": course.get("department_id"), "department_name": (dept or {}).get("name", ""), "teacher_id": course.get("teacher_id")},
        "group": {k: g[k] for k in ("key", "name", "teacher_id", "teacher_name", "count")},
        "all_groups": [{"key": x["key"], "name": x["name"], "count": x["count"]} for x in view["groups"]],
        "students": students,
        "limits": await course_limits(db, course),
        "schedule": sorted([{"day": c.get("day"), "slot_number": c.get("slot_number"), "start_time": c.get("start_time"), "end_time": c.get("end_time"), "room": rooms.get(c.get("room_id") or "", ""), "slot_type": c.get("slot_type")} for c in cells], key=lambda x: (str(x["day"]), x["slot_number"] or 0)),
        "lectures": {"total": sum(lec_stats.values()), **lec_stats},
    }


# ==================== تصدير ====================

@router.get("/courses/{course_id}/groups/export")
async def export_course_groups(course_id: str, group: Optional[str] = None, current_user: dict = Depends(get_current_user)):
    db = get_db()
    course = await _course(db, course_id)
    keys = await teacher_group_keys(db, course, current_user)
    if keys is not None and group and group not in keys:
        raise HTTPException(status_code=403, detail="هذه المجموعة ليست ضمن مجموعاتك")
    view = restrict_view(await _view(db, course), keys)
    wb = openpyxl.Workbook()
    groups = [g for g in view["groups"] if not group or g["key"] == group]
    if not groups and group:
        raise HTTPException(status_code=404, detail="المجموعة غير معرّفة")
    first = True
    for g in groups:
        _sheet(wb, g["name"], GROUP_HEADERS, [[s["student_number"], s["full_name"], g["name"], g["key"], g.get("teacher_name") or ""] for s in g["students"]], first=first)
        first = False
    if not group and keys is None:
        _sheet(wb, "بلا مجموعة", GROUP_HEADERS, [[s["student_number"], s["full_name"], "", "", ""] for s in view["unassigned"]], first=first)
        first = False
        _sheet(wb, "الكل", GROUP_HEADERS, _group_rows(view))
    if first:
        _sheet(wb, "الكل", GROUP_HEADERS, _group_rows(view), first=True)
    name = f"مجموعات {course.get('name', '')}" + (f" - {groups[0]['name']}" if group else "")
    return _xlsx_response(wb, name)


@router.get("/groups/export-level")
async def export_level_groups(department_id: str = Query(...), level: int = Query(...), section: Optional[str] = Query(None), current_user: dict = Depends(get_current_user)):
    """ملف واحد: ورقة لكل مقرر في المستوى له مجموعات"""
    db = get_db()
    q = {"department_id": department_id, "level": level, "is_active": {"$ne": False}, "groups.0": {"$exists": True}}
    if section:
        q["section"] = section
    courses = await db.courses.find(q).sort("name", 1).to_list(200)
    if not courses:
        raise HTTPException(status_code=404, detail="لا توجد مقررات بمجموعات في هذا المستوى")
    wb = openpyxl.Workbook()
    summary = []
    first = True
    for c in courses:
        view = await _view(db, c)
        _sheet(wb, c.get("name", "مقرر"), GROUP_HEADERS, _group_rows(view), first=first)
        first = False
        for g in view["groups"]:
            summary.append([c.get("name", ""), g["name"], g["key"], g.get("teacher_name") or "", g["count"]])
        summary.append([c.get("name", ""), "بلا مجموعة", "", "", len(view["unassigned"])])
    _sheet(wb, "ملخص", ["المقرر", "المجموعة", "الرمز", "القائد", "عدد الطلاب"], summary)
    wb.move_sheet("ملخص", offset=-(len(wb.sheetnames) - 1))
    dept = await db.departments.find_one({"_id": ObjectId(department_id)}, {"name": 1}) if ObjectId.is_valid(department_id) else None
    return _xlsx_response(wb, f"مجموعات {(dept or {}).get('name', '')} م{level}{' ' + section if section else ''}")


@router.get("/courses/{course_id}/groups/import-template")
async def import_template(course_id: str, current_user: dict = Depends(get_current_user)):
    db = get_db()
    course = await _course(db, course_id)
    view = await _view(db, course)
    wb = openpyxl.Workbook()
    ws = _sheet(wb, "التوزيع", ["رقم القيد", "اسم الطالب", "المجموعة"], [[r[0], r[1], r[3] or r[2]] for r in _group_rows(view)], first=True)
    ws["E1"], ws["E2"] = "المجموعات المعرّفة:", "اكتب رمز المجموعة أو اسمها في عمود «المجموعة» — فارغ = بلا مجموعة"
    for i, g in enumerate(view["groups"], 3):
        ws[f"E{i}"] = f"{g['key']} = {g['name']}"
    return _xlsx_response(wb, f"قالب توزيع {course.get('name', '')}")


# ==================== استيراد ====================

def _norm(v) -> str:
    s = str(v if v is not None else "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


async def _parse_import(db, course: dict, data: bytes) -> dict:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="تعذر قراءة ملف Excel")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="الملف فارغ")
    header = [_norm(h) for h in rows[0]]
    def col(*names):
        for n in names:
            for i, h in enumerate(header):
                if n in h:
                    return i
        return None
    c_id, c_grp = col("رقم القيد", "القيد", "student_id", "رقم"), col("المجموعة", "group", "مجموعة")
    if c_id is None or c_grp is None:
        raise HTTPException(status_code=400, detail="الأعمدة المطلوبة: «رقم القيد» و«المجموعة»")
    cid = str(course["_id"])
    enrollments = await db.enrollments.find({"course_id": cid}, {"student_id": 1, "group": 1}).to_list(10000)
    sids = [ObjectId(e["student_id"]) for e in enrollments if ObjectId.is_valid(e["student_id"])]
    students = {s.get("student_id", ""): s for s in await db.students.find({"_id": {"$in": sids}}, {"student_id": 1, "full_name": 1}).to_list(10000)}
    cur_group = {e["student_id"]: e.get("group") or "" for e in enrollments}
    groups = course.get("groups") or []
    gmap = {}
    for g in groups:
        gmap[_norm(g["key"]).lower()] = g["key"]
        gmap[_norm(g.get("name")).lower()] = g["key"]
    valid, unknown_students, undefined_groups, unchanged = [], [], {}, 0
    for r in rows[1:]:
        if r is None or len(r) <= max(c_id, c_grp):
            continue
        sno, graw = _norm(r[c_id]), _norm(r[c_grp])
        if not sno:
            continue
        st = students.get(sno)
        if not st:
            unknown_students.append({"student_number": sno, "group": graw})
            continue
        gkey = gmap.get(graw.lower()) if graw else ""
        if graw and gkey is None:
            undefined_groups.setdefault(graw, []).append(sno)
            gkey = graw  # يُنشأ إن طُلب
        sid = str(st["_id"])
        if (cur_group.get(sid) or "") == (gkey or ""):
            unchanged += 1
        valid.append({"student_db_id": sid, "student_number": sno, "full_name": st.get("full_name", ""), "group": gkey or "", "current_group": cur_group.get(sid) or ""})
    return {"rows": valid, "unknown_students": unknown_students, "undefined_groups": [{"label": k, "count": len(v)} for k, v in undefined_groups.items()],
            "unchanged": unchanged, "changes": len([v for v in valid if v["group"] != v["current_group"]]), "total_rows": len(rows) - 1}


@router.post("/courses/{course_id}/groups/import-preview")
async def import_preview(course_id: str, file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    if not _can(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    db = get_db()
    course = await _course(db, course_id)
    res = await _parse_import(db, course, await file.read())
    res["rows"] = res["rows"][:300]
    return res


@router.post("/courses/{course_id}/groups/import-apply")
async def import_apply(course_id: str, file: UploadFile = File(...), create_missing_groups: bool = Query(False), current_user: dict = Depends(get_current_user)):
    if not _can(current_user):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    db = get_db()
    course = await _course(db, course_id)
    res = await _parse_import(db, course, await file.read())
    if res["undefined_groups"]:
        if not create_missing_groups:
            raise HTTPException(status_code=400, detail="توجد مجموعات غير معرّفة: " + "، ".join(u["label"] for u in res["undefined_groups"]) + " — فعّل «إنشاء المجموعات الناقصة» أو صحّح الملف")
        groups = list(course.get("groups") or [])
        for u in res["undefined_groups"]:
            key = u["label"][:10]
            groups.append({"key": key, "name": f"مجموعة {key}" if len(key) <= 2 else key, "teacher_id": None})
        await db.courses.update_one({"_id": course["_id"]}, {"$set": {"groups": groups}})
    applied = 0
    for r in res["rows"]:
        if r["group"] == r["current_group"]:
            continue
        upd = {"$set": {"group": r["group"]}} if r["group"] else {"$unset": {"group": ""}}
        await db.enrollments.update_one({"course_id": course_id, "student_id": r["student_db_id"]}, upd)
        applied += 1
    await log_activity(current_user, "import_course_groups", "course", course_id, course.get("name", ""),
                       {"summary": f"استيراد توزيع المجموعات لمقرر «{course.get('name', '')}» — {applied} تغيير", "unknown": len(res["unknown_students"])})
    return {"message": f"تم تطبيق {applied} تغيير" + (f" — {len(res['unknown_students'])} رقم قيد غير مسجل تم تجاهله" if res["unknown_students"] else ""),
            "applied": applied, "unknown_students": res["unknown_students"], "created_groups": [u["label"] for u in res["undefined_groups"]] if create_missing_groups else []}


# ==================== خريطة مجموعات الطلاب (لقائمة الطلاب) ====================

class IdsIn(BaseModel):
    student_ids: List[str]


@router.post("/students/groups-map")
async def students_groups_map(data: IdsIn, current_user: dict = Depends(get_current_user)):
    db = get_db()
    ids = data.student_ids[:3000]
    enrollments = await db.enrollments.find({"student_id": {"$in": ids}, "group": {"$nin": [None, ""]}}, {"student_id": 1, "course_id": 1, "group": 1}).to_list(20000)
    cids = list({e["course_id"] for e in enrollments if ObjectId.is_valid(e["course_id"])})
    courses = {str(c["_id"]): c for c in await db.courses.find({"_id": {"$in": [ObjectId(c) for c in cids]}}, {"name": 1, "groups": 1}).to_list(2000)} if cids else {}
    out: dict = {}
    for e in enrollments:
        c = courses.get(e["course_id"])
        if not c:
            continue
        g = next((x for x in (c.get("groups") or []) if x.get("key") == e["group"]), None)
        out.setdefault(e["student_id"], []).append({"course_id": e["course_id"], "course_name": c.get("name", ""), "group": e["group"], "group_name": (g or {}).get("name") or f"مجموعة {e['group']}"})
    return out
