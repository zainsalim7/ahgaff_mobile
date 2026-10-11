"""📊 كشف الساعات الإضافية للأساتذة — امتداد لتقرير نصاب المدرسين
- سياسة لكل كلية: تفعيل الإضافة (ساعة لكل N ساعات، كسرياً)، المُقسِّم، سعر الساعة الإضافية حسب الرتبة الأكاديمية.
- الحساب لكل مقرر: الافتراضية = الأسبوعية × الأسابيع الكاملة من بداية المحاضرات (حد قاطع، الإجازات لا تُخصم بل تُحتسب ساعاتها للأستاذ)، المنجزة، النقص، نسبة الإنجاز، ملاحظات تلقائية.
"""
import io
from datetime import datetime, timedelta
from typing import Optional, List, Dict

from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from bson import ObjectId

from .deps import get_db, get_current_user, has_permission, export_filename, export_headers
from models.permissions import UserRole, Permission
from models.lectures import LectureStatus
from ._active_semester import get_active_semester, normalize_sem_date

router = APIRouter(tags=["التقارير - الساعات الإضافية"])

LOW_RATE_THRESHOLD = 75
ACADEMIC_TITLES = ["أستاذ", "أستاذ مشارك", "أستاذ مساعد", "محاضر", "معيد"]
DEFAULT_POLICY = {"bonus_enabled": False, "bonus_divisor": 6, "rates": {}, "currency": "ر.ي", "lecture_hour_minutes": 50}


class PolicyIn(BaseModel):
    faculty_id: Optional[str] = ""
    bonus_enabled: Optional[bool] = None
    bonus_divisor: Optional[float] = None
    rates: Optional[Dict[str, float]] = None
    currency: Optional[str] = None
    lecture_hour_minutes: Optional[int] = None


# ==================== السياسة ====================
async def load_policies(db) -> Dict[str, dict]:
    """{faculty_id: policy} مع "" للافتراضي العام"""
    out = {"": dict(DEFAULT_POLICY)}
    async for d in db.overtime_policies.find({}):
        fid = d.get("faculty_id") or ""
        out[fid] = {**DEFAULT_POLICY, **{k: v for k, v in d.items() if k in DEFAULT_POLICY}}
    return out


def policy_for(policies: Dict[str, dict], faculty_id: Optional[str]) -> dict:
    return policies.get(faculty_id or "") or policies[""]


def _can_manage_policy(user: dict, faculty_id: str) -> bool:
    if user.get("role") == UserRole.ADMIN or has_permission(user, Permission.MANAGE_SETTINGS):
        return True
    return bool(faculty_id) and user.get("faculty_id") == faculty_id and has_permission(user, Permission.REPORT_TEACHER_WORKLOAD)


@router.get("/reports/teacher-workload/policies")
async def list_policies(current_user: dict = Depends(get_current_user)):
    if current_user.get("role") == UserRole.TEACHER:
        raise HTTPException(status_code=403, detail="غير مصرح")
    db = get_db()
    pols = await load_policies(db)
    faculties = [{"id": str(f["_id"]), "name": f.get("name", "")} for f in await db.faculties.find({}, {"name": 1}).sort("name", 1).to_list(200)]
    items = [{"faculty_id": "", "faculty_name": "الافتراضي العام", **pols[""], "is_default": True}]
    for f in faculties:
        p = pols.get(f["id"])
        items.append({"faculty_id": f["id"], "faculty_name": f["name"], **(p or pols[""]), "inherits_default": p is None})
    return {"items": items, "titles": ACADEMIC_TITLES, "threshold": LOW_RATE_THRESHOLD}


@router.put("/reports/teacher-workload/policy")
async def put_policy(data: PolicyIn, current_user: dict = Depends(get_current_user)):
    fid = data.faculty_id or ""
    if not _can_manage_policy(current_user, fid):
        raise HTTPException(status_code=403, detail="غير مصرح لك بتعديل سياسة هذه الكلية")
    db = get_db()
    upd = {k: v for k, v in data.model_dump(exclude_none=True).items() if k != "faculty_id"}
    if "bonus_divisor" in upd and upd["bonus_divisor"] <= 0:
        raise HTTPException(status_code=400, detail="المُقسِّم يجب أن يكون أكبر من صفر")
    if "rates" in upd:
        upd["rates"] = {k: float(v) for k, v in upd["rates"].items() if v is not None and float(v) > 0}
    upd["updated_at"] = datetime.utcnow().isoformat()
    upd["updated_by"] = current_user.get("full_name", "")
    await db.overtime_policies.update_one({"faculty_id": fid}, {"$set": {**upd, "faculty_id": fid}}, upsert=True)
    return {"message": "تم حفظ سياسة الساعات الإضافية", **policy_for(await load_policies(db), fid)}


@router.delete("/reports/teacher-workload/policy/{faculty_id}")
async def delete_policy(faculty_id: str, current_user: dict = Depends(get_current_user)):
    """إعادة الكلية إلى الافتراضي العام"""
    if not _can_manage_policy(current_user, faculty_id):
        raise HTTPException(status_code=403, detail="غير مصرح")
    await get_db().overtime_policies.delete_one({"faculty_id": faculty_id})
    return {"message": "أُعيدت الكلية إلى السياسة الافتراضية"}


# ==================== الحساب ====================
def _lec_hours(lecture: dict) -> float:
    if lecture.get("credited_minutes"):
        return float(lecture["credited_minutes"]) / 60
    try:
        s = datetime.strptime(lecture.get("start_time", "00:00"), "%H:%M")
        e = datetime.strptime(lecture.get("end_time", "00:00"), "%H:%M")
        return ((e - s).seconds / 3600) or 1.0
    except Exception:
        return 1.0


def _d(v) -> Optional[datetime]:
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    try:
        return datetime.fromisoformat(str(v)[:10])
    except Exception:
        return None


def weeks_between(start: datetime, end: datetime) -> int:
    """عدد الأسابيع الكاملة (كل 7 أيام = أسبوع، أدناه 1) — حد قاطع كما في ترويسة الكشف"""
    days = (end.date() - start.date()).days + 1
    return max(1, days // 7)


AR_DAYS = {"الإثنين": 0, "الاثنين": 0, "الثلاثاء": 1, "الأربعاء": 2, "الخميس": 3, "الجمعة": 4, "السبت": 5, "الأحد": 6,
           "monday": 0, "tuesday": 1, "wednesday": 2, "thursday": 3, "friday": 4, "saturday": 5, "sunday": 6}


async def _course_day_hours(db, course_id: str, lectures: List[dict], weekly: float) -> Dict[int, float]:
    """{يوم الأسبوع: ساعات الجلسة} من الجدول الأسبوعي، وإلا من أيام المحاضرات الفعلية (توزيع متساوٍ)"""
    out: Dict[int, float] = {}
    async for s in db.weekly_schedule.find({"course_id": course_id}, {"day": 1, "day_of_week": 1, "start_time": 1, "end_time": 1}):
        d = s.get("day_of_week")
        if d is None:
            d = AR_DAYS.get(str(s.get("day", "")).strip().lower()) if str(s.get("day", "")).lower() in AR_DAYS else AR_DAYS.get(str(s.get("day", "")).strip())
        if d is not None:
            out[int(d)] = out.get(int(d), 0.0) + _lec_hours(s)
    if out:
        return out
    days = {d.weekday() for d in (_d(l.get("date")) for l in lectures) if d}
    return {d: round(weekly / len(days), 2) for d in days} if days else {}


def _holiday_deduction(holidays: Dict[str, str], day_hours: Dict[int, float], first: datetime, end: datetime, lecture_dates: set) -> (float, List[str]):
    """ساعات الإجازات الرسمية التي تصادف أيام المقرر بين أول محاضرة ونهاية الفترة (ما لم تُنفَّذ محاضرة فعلاً في ذلك اليوم)"""
    total, names = 0.0, []
    for ds, name in holidays.items():
        d = _d(ds)
        if not d or d < first or d > end or ds in lecture_dates:
            continue
        h = day_hours.get(d.weekday())
        if h:
            total += h
            if name not in names:
                names.append(name)
    return round(total, 2), names


async def _course_weekly_hours(db, teacher_id: str, course: dict) -> float:
    """الساعات الأسبوعية للمقرر: من إعدادات المقرر نفسه (الساعات المعتمدة/الأسبوعية) أولاً، ثم عبء التدريس، ثم مجموع الجدول الأسبوعي"""
    for k in ("weekly_hours", "credit_hours", "hours"):
        try:
            v = float(course.get(k) or 0)
        except (TypeError, ValueError):
            v = 0
        if v > 0:
            return v
    load = await db.teaching_loads.find_one({"teacher_id": teacher_id, "course_id": str(course["_id"])})
    if load and load.get("weekly_hours"):
        return float(load["weekly_hours"])
    total = 0.0
    async for sch in db.weekly_schedule.find({"course_id": str(course["_id"])}, {"start_time": 1, "end_time": 1}):
        total += _lec_hours(sch)
    return round(total, 2) if total else 3.0


async def semester_for_period(db, start_str: str, end_str: str) -> Optional[dict]:
    """الفصل الذي تقع فيه فترة التقرير حسب التاريخ (وإلا الفصل النشط)"""
    best = None
    async for sem in db.semesters.find({}):
        s0, e0 = normalize_sem_date(sem.get("start_date")), normalize_sem_date(sem.get("end_date"))
        if s0 and e0 and s0 <= end_str and e0 >= start_str:
            if best is None or (sem.get("status") == "active" or sem.get("is_active")):
                best = sem
    if best:
        return {"id": str(best["_id"]), "name": best.get("name", ""), "academic_year": best.get("academic_year", ""), "by_date": True}
    act = await get_active_semester(db)
    if act:
        act["by_date"] = False
    return act


def report_totals(result: list) -> dict:
    """إجماليات الكشف (تُعاد عند استبعاد أساتذة من الخطاب)"""
    return {
        "total_teachers": len(result),
        "total_required_hours": round(sum(t["summary"]["required_hours"] for t in result), 2),
        "total_scheduled_hours": round(sum(t["summary"]["total_scheduled_hours"] for t in result), 2),
        "total_actual_hours": round(sum(t["summary"]["total_actual_hours"] for t in result), 2),
        "total_difference_hours": round(sum(t["summary"]["difference_hours"] for t in result), 2),
        "total_expected_hours": round(sum(t["summary"]["expected_hours"] for t in result), 2),
        "total_shortfall_hours": round(sum(t["summary"]["shortfall_hours"] for t in result), 2),
        "total_bonus_hours": round(sum(t["summary"]["bonus_hours"] for t in result), 2),
        "total_overtime_hours": round(sum(max(0.0, t["summary"]["overtime_hours"]) for t in result), 2),
        "total_overtime_amount": round(sum(t["summary"]["overtime_amount"] or 0 for t in result), 2),
        "threshold": LOW_RATE_THRESHOLD,
    }


async def compute_workload(db, teachers: List[dict], start: datetime, end: datetime, hide_empty: bool = False, only_overtime: bool = False, exclude_ids: Optional[set] = None) -> dict:
    """يعيد نفس بنية /reports/teacher-workload القديمة + أعمدة كشف الساعات الإضافية"""
    from .calendar_events import holiday_dates
    policies = await load_policies(db)
    total_weeks = weeks_between(start, end)
    start_str, end_str = start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")
    sem = await semester_for_period(db, start_str, end_str)
    sem_id = (sem or {}).get("id")
    holidays = await holiday_dates(db, start_str, end_str)
    fac_names = {str(f["_id"]): f.get("name", "") for f in await db.faculties.find({}, {"name": 1}).to_list(300)}
    dep_docs = {str(d["_id"]): d for d in await db.departments.find({}, {"name": 1, "faculty_id": 1}).to_list(1000)}

    result = []
    for teacher in teachers:
        tid = str(teacher["_id"])
        weekly_hours = float(teacher.get("weekly_hours") or 12)
        required_hours = round(weekly_hours * total_weeks, 2)
        dep = dep_docs.get(teacher.get("department_id") or "", {})
        faculty_id = teacher.get("faculty_id") or dep.get("faculty_id") or ""
        pol = policy_for(policies, faculty_id)

        courses = await db.courses.find({"teacher_id": tid, "is_active": True}).to_list(100)
        if sem_id:
            # نطاق الفصل: مقررات فصل الفترة فقط (المقررات بلا فصل تبقى)
            courses = [c for c in courses if not c.get("semester_id") or str(c.get("semester_id")) == sem_id]
        cur_ids = {str(c["_id"]) for c in courses}
        for hc in await db.lectures.distinct("course_id", {"teacher_id": tid, "date": {"$gte": start_str, "$lte": end_str}}):
            if hc not in cur_ids and ObjectId.is_valid(hc):
                cdoc = await db.courses.find_one({"_id": ObjectId(hc)})
                if cdoc:
                    cdoc["_hist_only"] = True
                    courses.append(cdoc)

        total_scheduled = total_actual = total_expected = total_shortfall = total_excused = 0.0
        courses_data = []
        for course in courses:
            cid = str(course["_id"])
            attr = [{"teacher_id": tid}] if course.get("_hist_only") else [{"teacher_id": tid}, {"teacher_id": None}]
            lectures = await db.lectures.find({
                "course_id": cid,
                "$and": [{"$or": [{"date": {"$gte": start_str, "$lte": end_str}}, {"date": {"$gte": start, "$lte": end}}]}, {"$or": attr}],
                "is_cancelled": {"$ne": True}, "status": {"$ne": LectureStatus.CANCELLED},
            }).to_list(1000)
            lec_ids = [str(l["_id"]) for l in lectures]
            executed_ids = set(await db.attendance.distinct("lecture_id", {"lecture_id": {"$in": lec_ids}})) if lec_ids else set()
            # محاضرات تقع في إجازة رسمية ولم تُنفَّذ تُستبعد من المجدولة (لا تُعد نقصاً)
            executed_dates = {str(l.get("date"))[:10] for l in lectures if str(l["_id"]) in executed_ids}
            lectures = [l for l in lectures if str(l.get("date"))[:10] not in holidays or str(l["_id"]) in executed_ids]
            excused = [l for l in lectures if l.get("status") == LectureStatus.ABSENT_EXCUSED]
            executed = [l for l in lectures if str(l["_id"]) in executed_ids]

            scheduled_hours = round(sum(_lec_hours(l) for l in lectures), 2)
            actual_hours = round(sum(_lec_hours(l) for l in executed), 2)
            excused_hours = round(sum(_lec_hours(l) for l in excused), 2)
            weekly = await _course_weekly_hours(db, tid, course)

            # بداية المقرر = تاريخ بداية توليد المحاضرات (أول محاضرة مجدولة في الفترة)
            dates = [d for d in (_d(l.get("date")) for l in lectures) if d]
            first = min(dates) if dates else None
            notes = []
            start_week = 1
            if first:
                course_weeks = weeks_between(first, end)
                start_week = (first.date() - start.date()).days // 7 + 1
                if start_week > 1:
                    notes.append(f"بدأ من الأسبوع {start_week} (بداية المحاضرات {first.strftime('%Y-%m-%d')}؛ الافتراضية من {course_weeks} أسابيع)")
            else:
                course_weeks = 0
                notes.append("لا محاضرات مجدولة في الفترة")
            # الإجازات الرسمية لا تُخصم من الافتراضية ولا النصاب؛ ساعاتها تُحتسب للمدرس كمنفَّذة (لا يظهر نقص)
            holiday_hours, holiday_names = 0.0, []
            if first and holidays:
                day_hours = await _course_day_hours(db, cid, lectures, weekly)
                holiday_hours, holiday_names = _holiday_deduction(holidays, day_hours, first, end, executed_dates)
            taught_hours = actual_hours
            actual_hours = round(actual_hours + holiday_hours, 2)
            expected = round(max(0.0, weekly * course_weeks - excused_hours), 2)
            if excused_hours:
                notes.append(f"غياب بعذر {excused_hours:g} س (خُصم من الافتراضية)")
            if holiday_hours:
                notes.append(f"إجازة رسمية {holiday_hours:g} س محتسبة ({'، '.join(holiday_names)})")
            shortfall = round(max(0.0, expected - actual_hours), 2)
            rate = round(actual_hours / expected * 100) if expected > 0 else (100 if actual_hours > 0 else 0)

            total_scheduled += scheduled_hours; total_actual += actual_hours
            total_expected += expected; total_shortfall += shortfall; total_excused += excused_hours
            courses_data.append({
                "course_id": cid, "course_name": course.get("name", ""), "course_code": course.get("code", ""),
                "department_id": course.get("department_id"), "department_name": dep_docs.get(course.get("department_id") or "", {}).get("name", ""),
                "level": course.get("level"), "section": course.get("section") or "",
                "weekly_hours": weekly, "first_lecture_date": first.strftime("%Y-%m-%d") if first else None, "course_weeks": course_weeks,
                "start_week": start_week, "late_start": start_week > 1,
                "scheduled_lectures": len(lectures), "executed_lectures": len(executed), "excused_lectures": len(excused),
                "scheduled_hours": scheduled_hours, "actual_hours": actual_hours, "taught_hours": taught_hours, "excused_hours": excused_hours,
                "expected_hours": expected, "holiday_hours": holiday_hours, "shortfall_hours": shortfall, "completion_rate": rate, "is_low": expected > 0 and rate < LOW_RATE_THRESHOLD,
                "note": "، ".join(notes),
            })

        courses_data.sort(key=lambda c: ((c["level"] or 0), c["course_name"], c["section"]))
        # نصاب المدرس حد قاطع: الأسبوعي × الأسابيع الكاملة؛ يبدأ من بداية محاضراته إن كان توليدها بعد بداية الفترة
        firsts = [_d(c["first_lecture_date"]) for c in courses_data if c.get("first_lecture_date")]
        teacher_first = min(firsts) if firsts else None
        effective_weeks, teacher_start_week, teacher_note = total_weeks, 1, ""
        if teacher_first:
            teacher_start_week = (teacher_first.date() - start.date()).days // 7 + 1
            if teacher_start_week > 1:
                effective_weeks = weeks_between(teacher_first, end)
                required_hours = round(weekly_hours * effective_weeks, 2)
                teacher_note = (f"بداية المحاضرات {teacher_first.strftime('%Y-%m-%d')} (الأسبوع {teacher_start_week})"
                                f" · النصاب محسوب من {effective_weeks} أسابيع بدل {total_weeks}")
        bonus_enabled = bool(pol.get("bonus_enabled"))
        divisor = float(pol.get("bonus_divisor") or 6)
        bonus_hours = round(total_actual / divisor, 2) if bonus_enabled else 0.0
        after_bonus = round(total_actual + bonus_hours, 2)
        overtime = round(after_bonus - required_hours, 2)
        title = teacher.get("academic_title") or ""
        rate_val = (pol.get("rates") or {}).get(title)
        amount = round(max(0.0, overtime) * float(rate_val), 2) if rate_val else None
        result.append({
            "teacher_id": teacher.get("teacher_id", "") or tid, "teacher_db_id": tid,
            "teacher_name": teacher.get("full_name", ""), "academic_title": title,
            "department_id": teacher.get("department_id"), "department_name": dep.get("name", ""),
            "faculty_id": faculty_id, "faculty_name": fac_names.get(faculty_id, ""),
            "weekly_hours": weekly_hours, "courses": courses_data,
            "summary": {
                "total_courses": len(courses_data), "weekly_hours": weekly_hours, "total_weeks": total_weeks,
                "effective_weeks": effective_weeks, "start_week": teacher_start_week, "late_start": teacher_start_week > 1,
                "first_lecture_date": teacher_first.strftime("%Y-%m-%d") if teacher_first else None, "note": teacher_note,
                "holiday_hours": round(sum(c["holiday_hours"] for c in courses_data), 2),
                "required_hours": required_hours, "total_scheduled_hours": round(total_scheduled, 2), "total_actual_hours": round(total_actual, 2),
                "difference_hours": round(total_actual - required_hours, 2),
                "completion_rate": round(total_actual / required_hours * 100, 2) if required_hours > 0 else 0,
                "expected_hours": round(total_expected, 2), "shortfall_hours": round(total_shortfall, 2), "excused_hours": round(total_excused, 2),
                "bonus_enabled": bonus_enabled, "bonus_divisor": divisor, "bonus_hours": bonus_hours,
                "hours_after_bonus": after_bonus, "overtime_hours": overtime,
                "hourly_rate": rate_val, "currency": pol.get("currency", ""), "overtime_amount": amount,
            },
        })

    if hide_empty:
        result = [t for t in result if t["summary"]["total_scheduled_hours"] > 0]
    if only_overtime:
        result = [t for t in result if t["summary"]["overtime_hours"] > 0]
    if exclude_ids:
        result = [t for t in result if t["teacher_db_id"] not in exclude_ids]
    return {
        "holidays": [{"date": k, "name": v} for k, v in sorted(holidays.items())],
        "period": {"start_date": start.isoformat(), "end_date": end.isoformat(), "total_weeks": total_weeks, "holiday_days": len(holidays),
                   "semester_name": (sem or {}).get("name", ""), "academic_year": (sem or {}).get("academic_year", "")},
        "teachers": result,
        "summary": report_totals(result),
    }


async def select_teachers(db, current_user: dict, teacher_id: Optional[str], faculty_id: Optional[str], department_id: Optional[str]) -> List[dict]:
    """اختيار المدرسين حسب الدور والفلاتر (المدرس: نفسه فقط)"""
    if current_user["role"] == UserRole.TEACHER:
        rec = await db.teachers.find_one({"user_id": current_user["id"]})
        if not rec:
            raise HTTPException(status_code=404, detail="لم يتم العثور على سجل المعلم")
        return [rec]
    if not has_permission(current_user, Permission.VIEW_REPORTS) and not has_permission(current_user, Permission.REPORT_TEACHER_WORKLOAD):
        raise HTTPException(status_code=403, detail="غير مصرح لك")
    q: dict = {"is_active": {"$ne": False}}
    if teacher_id and ObjectId.is_valid(teacher_id):
        q["_id"] = ObjectId(teacher_id)
    if department_id:
        q["$or"] = [{"department_id": department_id}, {"department_ids": department_id}]
    elif faculty_id:
        dep_ids = [str(d["_id"]) for d in await db.departments.find({"faculty_id": faculty_id}, {"_id": 1}).to_list(500)]
        q["$or"] = [{"faculty_id": faculty_id}, {"department_id": {"$in": dep_ids}}, {"department_ids": {"$in": dep_ids}}]
    return await db.teachers.find(q).sort("full_name", 1).to_list(1000)


def parse_period(start_date: Optional[str], end_date: Optional[str]):
    from routes.notifications import get_yemen_time
    if not start_date or not end_date:
        today = get_yemen_time().replace(tzinfo=None)
        start = today.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        end = (start + timedelta(days=32)).replace(day=1) - timedelta(seconds=1)
    else:
        start = datetime.fromisoformat(start_date)
        end = datetime.fromisoformat(end_date)
    return start, end


# ==================== الكشف: PDF أفقي + Excel ====================
SHEET_HEADERS = ["م", "اسم الأستاذ", "المادة", "المستوى", "الساعات الأسبوعية", "النصاب الأسبوعي", "النصاب", "الساعات الافتراضية",
                 "الساعات المنجزة", "إجمالي المنجز", "إضافة ساعة لكل {n} ساعات", "الساعات بعد الإضافة ({m} دقيقة)",
                 "إجمالي الساعات الإضافية", "النقص", "نسبة الإنجاز", "الملاحظات"]


BONUS_COLS = (10, 11)  # «إضافة ساعة لكل N» و«الساعات بعد الإضافة» — تظهر فقط إن فُعِّلت الإضافة لكلية ما


def any_bonus(report: dict) -> bool:
    return any(t["summary"].get("bonus_enabled") for t in report["teachers"])


def _cols(seq, with_bonus: bool):
    """يحذف عمودي الإضافة من قائمة بترتيب الأعمدة إن كانت الإضافة معطّلة"""
    return list(seq) if with_bonus else [v for i, v in enumerate(seq) if i not in BONUS_COLS]


def _lvl(c: dict) -> str:
    lv = {1: "الأول", 2: "الثاني", 3: "الثالث", 4: "الرابع", 5: "الخامس", 6: "السادس"}.get(c.get("level"), str(c.get("level") or ""))
    return f"{lv} {c.get('section') or ''}".strip()


def _fmt(v) -> str:
    if v is None or v == "":
        return ""
    try:
        return f"{float(v):.1f}"
    except Exception:
        return str(v)


def sheet_rows(report: dict):
    """صفوف الكشف: سطر لكل مقرر، وقيم الأستاذ تُكرَّر (تُدمج لاحقاً)؛ عمودا الإضافة يُحذفان إن كانت معطّلة"""
    rows, spans = [], []
    with_bonus = any_bonus(report)
    idx = 0
    for t in report["teachers"]:
        s = t["summary"]
        cs = t["courses"] or [{"course_name": "—", "level": None, "section": "", "weekly_hours": "", "expected_hours": "", "actual_hours": "", "shortfall_hours": "", "completion_rate": "", "note": "لا مقررات", "is_low": False}]
        idx += 1
        name = f"{t.get('academic_title') or ''} {t['teacher_name']}".strip()
        first = len(rows)
        for k, c in enumerate(cs):
            note = c.get("note", "")
            if k == 0 and s.get("note"):
                note = f"{s['note']}" + (f" | {note}" if note else "")
            rows.append({
                "cells": [idx, name, (f"{c['course_name']}\n{c['department_name']}" if c.get("department_name") else c["course_name"]), _lvl(c), _fmt(c.get("weekly_hours")), _fmt(s["weekly_hours"]), _fmt(s["required_hours"]), _fmt(c.get("expected_hours")),
                          _fmt(c.get("actual_hours")), _fmt(s["total_actual_hours"]), _fmt(s["bonus_hours"]) if s["bonus_enabled"] else "—", _fmt(s["hours_after_bonus"]),
                          _fmt(s["overtime_hours"]), _fmt(c.get("shortfall_hours")), f"{c['completion_rate']}%" if c.get("completion_rate") != "" else "", note],
                "is_low": bool(c.get("is_low")),
            })
            rows[-1]["cells"] = _cols(rows[-1]["cells"], with_bonus)
        spans.append((first, len(rows) - 1))
    return rows, spans


def build_sheet_pdf(report: dict, issuer: str, generated_by: str) -> io.BytesIO:
    import os
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import arabic_reshaper
    from bidi.algorithm import get_display

    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    font, bold = "Helvetica", "Helvetica-Bold"
    try:
        pdfmetrics.registerFont(TTFont("Amiri", os.path.join(here, "fonts", "Amiri-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("Amiri-Bold", os.path.join(here, "fonts", "Amiri-Bold.ttf")))
        font, bold = "Amiri", "Amiri-Bold"
    except Exception:
        pass

    def ar(t):
        try:
            return get_display(arabic_reshaper.reshape(str(t if t is not None else "")))
        except Exception:
            return str(t or "")

    def wrap_ar(t, maxc):
        """لفّ النص بالكلمات قبل تطبيق bidi حتى لا يختل ترتيب الكلمات عند تعدد الأسطر (\n = سطر جديد إجباري)"""
        if "\n" in str(t or ""):
            return "<br/>".join(wrap_ar(part, maxc) for part in str(t).split("\n"))
        words, lines, cur = str(t or "").split(), [], ""
        for w in words:
            if cur and len(cur) + 1 + len(w) > maxc:
                lines.append(cur); cur = w
            else:
                cur = f"{cur} {w}".strip()
        if cur:
            lines.append(cur)
        return "<br/>".join(ar(l) for l in lines) or ""

    NAVY, GREY, PINK = colors.HexColor("#0f2440"), colors.HexColor("#64748b"), colors.HexColor("#fbcfe8")
    st_title = ParagraphStyle("t", fontName=bold, fontSize=14, leading=20, alignment=TA_CENTER, textColor=NAVY)
    st_sub = ParagraphStyle("s", fontName=font, fontSize=9.5, leading=14, alignment=TA_CENTER, textColor=GREY)
    st_small = ParagraphStyle("m", fontName=font, fontSize=8, alignment=TA_RIGHT, textColor=GREY, leading=12)
    st_cell = ParagraphStyle("c", fontName=font, fontSize=6.6, leading=8.4, alignment=TA_CENTER)

    p = report["period"]
    pols = {t["summary"]["bonus_divisor"] for t in report["teachers"] if t["summary"]["bonus_enabled"]}
    div_label = f"{int(min(pols)) if pols and min(pols).is_integer() else (min(pols) if pols else 6)}"
    with_bonus = any_bonus(report)
    headers = _cols([h.format(n=div_label, m=50) for h in SHEET_HEADERS], with_bonus)
    rows, spans = sheet_rows(report)

    el = []
    logo = os.path.join(here, "assets", "university_logo.jpeg")
    if os.path.exists(logo):
        el.append(RLImage(logo, width=16 * mm, height=16 * mm))
    sem = f"للفصل {p.get('semester_name') or ''} {p.get('academic_year') or ''}".strip()
    el += [Paragraph(ar(f"كشف يبين الساعات الإضافية للأساتذة {sem}"), st_title),
           Paragraph(ar(f"الفترة من {str(p['start_date'])[:10]} إلى {str(p['end_date'])[:10]}   ·   عدد الأسابيع: {p['total_weeks']}   ·   أيام الإجازات الرسمية: {p.get('holiday_days', 0)}   ·   جهة الإصدار: {issuer or 'قسم التسجيل'}" + ("" if with_bonus else "   ·   الإضافة غير مفعّلة")), st_sub),
           Spacer(1, 3 * mm)]

    widths = _cols([7, 34, 30, 17, 13, 13, 12, 15, 14, 14, 16, 18, 16, 11, 12, 35], with_bonus)
    total_w = sum(widths)
    avail = landscape(A4)[0] - 16 * mm
    widths = [w / total_w * avail for w in widths]
    maxc = [max(4, int(w / mm * 1.3)) for w in widths]
    data = [[Paragraph(wrap_ar(h, maxc[i]), ParagraphStyle("h", parent=st_cell, fontName=bold, textColor=colors.white)) for i, h in reversed(list(enumerate(headers)))]]
    for r in rows:
        data.append([Paragraph(wrap_ar(c, maxc[i]), st_cell) for i, c in reversed(list(enumerate(r["cells"])))])
    t = Table(data, colWidths=list(reversed(widths)), repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), NAVY), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94a3b8")),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    ncols = len(headers)
    merge_cols = {0, 1, 5, 6, 9, 10, 11, 12} if with_bonus else {0, 1, 5, 6, 9, 10}  # أعمدة الأستاذ تُدمج عبر أسطره
    low_col = 14 if with_bonus else 12
    SEP = colors.HexColor("#334155")
    for (a, b) in spans:
        if b > a:
            for ci in merge_cols:
                col = ncols - 1 - ci
                style.append(("SPAN", (col, a + 1), (col, b + 1)))
        style.append(("LINEBELOW", (0, b + 1), (-1, b + 1), 1.3, SEP))  # فاصل أغمق بين أستاذ وآخر
    for i, r in enumerate(rows, start=1):
        if r["is_low"]:
            col = ncols - 1 - low_col
            style.append(("BACKGROUND", (col, i), (col, i), PINK))
    t.setStyle(TableStyle(style))
    el.append(t)
    s = report["summary"]
    el += [Spacer(1, 4 * mm),
           Paragraph(ar(f"الإجمالي: {s['total_teachers']} أستاذ · النصاب {s['total_required_hours']:g} س · المنجز {s['total_actual_hours']:g} س · الإضافة {s['total_bonus_hours']:g} س · الساعات الإضافية {s['total_overtime_hours']:g} س · النقص {s['total_shortfall_hours']:g} س"), st_small),
           Paragraph(wrap_ar("المنجز = محاضرات سُجّل لها حضور بمدتها الفعلية. الافتراضية = الساعات الأسبوعية × الأسابيع الكاملة من بداية المحاضرات حتى نهاية الفترة (يُخصم غياب الأستاذ بعذر فقط). النصاب = النصاب الأسبوعي × الأسابيع الكاملة (حد قاطع). الإجازات الرسمية لا تُخصم وتُحتسب ساعاتها للأستاذ. " + ("الإضافة = المنجز ÷ المُقسِّم (للكليات المفعَّل فيها). الساعات الإضافية = (المنجز + الإضافة) − النصاب. " if with_bonus else "الساعات الإضافية = المنجز − النصاب. ") + "تُظلَّل نسبة الإنجاز الأقل من 75٪."
                        + (f"  |  أُصدر بواسطة: {generated_by}" if generated_by else "") + f"  |  {datetime.now().strftime('%Y-%m-%d %H:%M')}", 150), st_small)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=8 * mm, leftMargin=8 * mm, topMargin=8 * mm, bottomMargin=8 * mm, title="كشف الساعات الإضافية").build(el)
    buf.seek(0)
    return buf


def build_sheet_excel(report: dict, issuer: str) -> io.BytesIO:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "كشف الساعات الإضافية"
    ws.sheet_view.rightToLeft = True
    p = report["period"]
    pols = {t["summary"]["bonus_divisor"] for t in report["teachers"] if t["summary"]["bonus_enabled"]}
    with_bonus = any_bonus(report)
    headers = _cols([h.format(n=f"{min(pols):g}" if pols else "6", m=50) for h in SHEET_HEADERS], with_bonus)
    n = len(headers)
    thin = Side(style="thin", color="94A3B8"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    c = ws.cell(row=1, column=1, value=f"كشف يبين الساعات الإضافية للأساتذة للفصل {p.get('semester_name') or ''} {p.get('academic_year') or ''}".strip()); c.font = Font(bold=True, size=14, color="0F2440"); c.alignment = center
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    c = ws.cell(row=2, column=1, value=f"الفترة من {str(p['start_date'])[:10]} إلى {str(p['end_date'])[:10]} · عدد الأسابيع: {p['total_weeks']} · أيام الإجازات الرسمية: {p.get('holiday_days', 0)} · جهة الإصدار: {issuer or 'قسم التسجيل'}"); c.font = Font(size=10, color="64748B"); c.alignment = center
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=i, value=h); c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = PatternFill(start_color="0F2440", end_color="0F2440", fill_type="solid"); c.alignment = center; c.border = border
    ws.row_dimensions[4].height = 42
    rows, spans = sheet_rows(report)
    pink = PatternFill(start_color="FBCFE8", end_color="FBCFE8", fill_type="solid")
    for ri, r in enumerate(rows, start=5):
        for ci, v in enumerate(r["cells"], 1):
            try:
                v = float(v) if isinstance(v, str) and v.replace(".", "", 1).isdigit() else v
            except Exception:
                pass
            c = ws.cell(row=ri, column=ci, value=v); c.alignment = center; c.border = border
        if r["is_low"]:
            ws.cell(row=ri, column=15 if with_bonus else 13).fill = pink
    thick = Side(style="medium", color="334155")
    for (a, b) in spans:
        if b > a:
            for ci in ((1, 2, 6, 7, 10, 11, 12, 13) if with_bonus else (1, 2, 6, 7, 10, 11)):
                ws.merge_cells(start_row=a + 5, start_column=ci, end_row=b + 5, end_column=ci)
        for ci in range(1, n + 1):
            c = ws.cell(row=b + 5, column=ci); c.border = Border(left=thin, right=thin, top=thin, bottom=thick)
    for i, w in enumerate(_cols([4, 26, 24, 14, 10, 10, 9, 12, 11, 11, 13, 14, 13, 9, 10, 30], with_bonus), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # ورقة المالية إن وُجد سعر ساعة
    fin = [t for t in report["teachers"] if t["summary"].get("hourly_rate")]
    if fin:
        ws2 = wb.create_sheet("المستحقات المالية")
        ws2.sheet_view.rightToLeft = True
        hs = ["#", "الأستاذ", "الرتبة", "الكلية", "القسم", "النصاب", "المنجز", "الإضافة", "الساعات الإضافية", "سعر الساعة", "المستحق", "العملة"]
        for i, h in enumerate(hs, 1):
            c = ws2.cell(row=1, column=i, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid"); c.alignment = center; c.border = border
        tot = 0.0
        for i, t in enumerate(fin, 1):
            s = t["summary"]; tot += s["overtime_amount"] or 0
            ws2.append([i, t["teacher_name"], t.get("academic_title", ""), t.get("faculty_name", ""), t.get("department_name", ""), s["required_hours"], s["total_actual_hours"], s["bonus_hours"], max(0.0, s["overtime_hours"]), s["hourly_rate"], s["overtime_amount"], s.get("currency", "")])
        ws2.append(["", "الإجمالي", "", "", "", "", "", "", "", "", round(tot, 2), ""])
        ws2.cell(row=ws2.max_row, column=11).font = Font(bold=True)
        for i, w in enumerate([4, 26, 14, 20, 20, 9, 9, 9, 12, 11, 14, 8], 1):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return out


async def _issuer(db, current_user: dict) -> str:
    if current_user.get("department_id") and ObjectId.is_valid(current_user["department_id"]):
        d = await db.departments.find_one({"_id": ObjectId(current_user["department_id"])}, {"name": 1})
        if d:
            return d.get("name", "")
    if current_user.get("faculty_id") and ObjectId.is_valid(current_user["faculty_id"]):
        f = await db.faculties.find_one({"_id": ObjectId(current_user["faculty_id"])}, {"name": 1})
        if f:
            return f.get("name", "")
    return "قسم التسجيل"


@router.get("/export/report/teacher-workload/sheet-pdf")
async def export_sheet_pdf(teacher_id: Optional[str] = None, start_date: Optional[str] = None, end_date: Optional[str] = None,
                           faculty_id: Optional[str] = None, department_id: Optional[str] = None, hide_empty: bool = False, only_overtime: bool = False, current_user: dict = Depends(get_current_user)):
    """📄 كشف الساعات الإضافية — PDF أفقي بشكل الكشف الرسمي"""
    db = get_db()
    start, end = parse_period(start_date, end_date)
    report = await compute_workload(db, await select_teachers(db, current_user, teacher_id, faculty_id, department_id), start, end, hide_empty, only_overtime)
    buf = build_sheet_pdf(report, await _issuer(db, current_user), current_user.get("full_name", ""))
    fname = export_filename("كشف الساعات الإضافية للأساتذة", f"{start.strftime('%Y-%m-%d')} الى {end.strftime('%Y-%m-%d')}", ext="pdf")
    return StreamingResponse(buf, media_type="application/pdf", headers=export_headers(fname))


@router.get("/export/report/teacher-workload/sheet-excel")
async def export_sheet_excel(teacher_id: Optional[str] = None, start_date: Optional[str] = None, end_date: Optional[str] = None,
                             faculty_id: Optional[str] = None, department_id: Optional[str] = None, hide_empty: bool = False, only_overtime: bool = False, current_user: dict = Depends(get_current_user)):
    """📊 كشف الساعات الإضافية — Excel (+ ورقة المستحقات المالية عند وجود سعر ساعة)"""
    db = get_db()
    start, end = parse_period(start_date, end_date)
    report = await compute_workload(db, await select_teachers(db, current_user, teacher_id, faculty_id, department_id), start, end, hide_empty, only_overtime)
    out = build_sheet_excel(report, await _issuer(db, current_user))
    fname = export_filename("كشف الساعات الإضافية للأساتذة", f"{start.strftime('%Y-%m-%d')} الى {end.strftime('%Y-%m-%d')}", ext="xlsx")
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=export_headers(fname))
